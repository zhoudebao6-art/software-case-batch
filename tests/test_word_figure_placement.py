import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from verify_docx import verify


class WordFigurePlacementTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.figure = self.root / 'figure.png'
        Image.new('RGB', (20, 20), 'gray').save(self.figure)
        self.original = self.root / 'original.docx'
        doc = Document()
        doc.add_paragraph('实施例前文')
        doc.add_paragraph('实施例后文')
        doc.save(self.original)

    def tearDown(self):
        self.tmp.cleanup()

    def check(self, doc):
        output = self.root / 'revised.docx'
        doc.save(output)
        return verify(self.original, output, [self.figure])

    def figure_block(self, doc):
        doc.add_picture(str(self.figure))
        doc.add_paragraph('图4 计算结果')

    def test_inline_figure_is_rejected_even_when_original_text_and_image_are_preserved(self):
        doc = Document()
        doc.add_paragraph('实施例前文')
        doc.add_paragraph('请参考图4，本实施例的结果说明。')
        self.figure_block(doc)
        doc.add_paragraph('实施例后文')
        result = self.check(doc)
        self.assertTrue(result['structural_preservation_ok'])
        self.assertTrue(result['all_requested_figures_embedded'])
        self.assertFalse(result['all_requested_figures_at_document_end'])

    def test_explanation_stays_in_body_and_figure_caption_go_to_end(self):
        doc = Document()
        doc.add_paragraph('实施例前文')
        doc.add_paragraph('请参考图4，本实施例的结果说明。')
        doc.add_paragraph('实施例后文')
        self.figure_block(doc)
        result = self.check(doc)
        self.assertTrue(result['structural_preservation_ok'])
        self.assertTrue(result['all_requested_figures_at_document_end'])
        self.assertTrue(result['all_requested_figures_referenced_in_body'])
        self.assertEqual(result['figure_placement']['captions'], ['图4 计算结果'])

    def test_caption_alone_does_not_replace_reference_in_body(self):
        doc = Document(self.original)
        self.figure_block(doc)
        result = self.check(doc)
        self.assertTrue(result['all_requested_figures_at_document_end'])
        self.assertFalse(result['all_requested_figures_referenced_in_body'])
        self.assertEqual(result['figure_placement']['missing_body_references'], [4])

    def test_wrong_figure_number_does_not_satisfy_reference(self):
        doc = Document(self.original)
        doc.add_paragraph('请参考图40，分析本实施例。')
        self.figure_block(doc)
        self.assertFalse(self.check(doc)['all_requested_figures_referenced_in_body'])

    def test_moving_explanation_with_figure_to_end_is_rejected(self):
        doc = Document(self.original)
        self.figure_block(doc)
        doc.add_paragraph('不应把结果说明移到图后。')
        self.assertFalse(self.check(doc)['all_requested_figures_at_document_end'])

    def test_unused_embedded_image_is_not_a_placed_figure(self):
        doc = Document(self.original)
        doc.part.get_or_add_image(str(self.figure))
        result = self.check(doc)
        self.assertTrue(result['all_requested_figures_embedded'])
        self.assertFalse(result['all_requested_figures_at_document_end'])

    def test_parent_rejects_inline_figure_before_rendering(self):
        import json
        from controlled_evidence import produce
        doc = Document()
        doc.add_paragraph('实施例前文')
        self.figure_block(doc)
        doc.add_paragraph('实施例后文')
        doc.save(self.root / 'revised.docx')
        evidence = self.root / 'evidence'
        evidence.mkdir()
        (evidence / 'artifact-manifest.json').write_text(json.dumps({
            'revised_docx': 'revised.docx', 'charts': [{'png': 'figure.png'}]
        }), encoding='utf-8')
        with patch('controlled_evidence.config_at', return_value={}), patch('controlled_evidence.subprocess.run') as render:
            with self.assertRaisesRegex(ValueError, 'must be at document end'):
                produce(self.root, self.original, 'word', self.root / 'config.json')
            render.assert_not_called()

    def test_duplicate_inline_and_end_figure_is_rejected(self):
        doc = Document()
        doc.add_paragraph('实施例前文')
        self.figure_block(doc)
        doc.add_paragraph('实施例后文')
        self.figure_block(doc)
        self.assertFalse(self.check(doc)['all_requested_figures_at_document_end'])

    def test_caption_above_figure_is_rejected(self):
        doc = Document(self.original)
        doc.add_paragraph('图4 计算结果')
        doc.add_picture(str(self.figure))
        self.assertFalse(self.check(doc)['all_requested_figures_at_document_end'])

    def test_original_figure_keeps_its_position(self):
        source = Document(self.original)
        original_image = self.root / 'source.png'
        Image.new('RGB', (20, 20), 'white').save(original_image)
        source.add_picture(str(original_image))
        source.add_paragraph('图3 原文流程图')
        source.save(self.original)
        self.figure_block(source)
        result = self.check(source)
        self.assertTrue(result['structural_preservation_ok'])
        self.assertTrue(result['all_requested_figures_at_document_end'])


if __name__ == '__main__':
    unittest.main()
